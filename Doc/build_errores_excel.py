"""Genera TCM_errores.xlsx — catálogo completo de errores (filas en rojo)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "TCM_errores.xlsx"

HEADERS = [
    "Módulo",
    "Categoría",
    "Código UI",
    "Tag / ID",
    "Byte (hex)",
    "Byte (dec)",
    "Wire L",
    "Wire R",
    "Lado",
    "Nombre / Descripción",
    "Función / handler",
    "Nivel",
    "Sistema / subsistema",
    "Fuente",
    "Es error",
]

# (module, category, ui_code, tag, hex, dec, wire_l, wire_r, side, name, handler, level, system, source)
ROWS: list[tuple] = []

# --- Catálogo unificado UI (Temp/Index.html ERROR_CATALOG) ---
UI_CATALOG = [
    (1, "Sensores", "Pinzas"),
    (2, "Sensores", "Sujetador"),
    (3, "Sensores", "Cortador"),
    (4, "Sensores", "Manguera (safety CAN)"),
    (5, "Sensores", "Bandeja"),
    (6, "Interlock", "Parada externa PreFeeder"),
    (7, "Sensores", "Sin comunicación CAN"),
    (8, "Servo", "Servo CAN no listo"),
    (10, "Lineal", "Timeout movimiento"),
    (11, "Alimentación", "Timeout feed (reserva)"),
    (12, "Alimentación", "Alimentación fallida (firmware E012, no en UI)"),
    (16, "Medición OM", "Reservado (láser)"),
    (17, "Medición OM", "Fase 1 fuera de tolerancia"),
    (18, "Medición OM", "Fase 2 fuera de tolerancia"),
    (19, "Medición OM", "Longitud total fuera de tolerancia"),
    (20, "Ciclo TCM", "Parada por seguridad"),
    (21, "Ciclo TCM", "Safety stop · pinzas"),
    (22, "Ciclo TCM", "Safety stop · sujetador"),
    (23, "Ciclo TCM", "Safety stop · cortador"),
    (24, "Ciclo TCM", "Safety stop · manguera"),
    (25, "Ciclo TCM", "Safety stop · bandeja"),
    (26, "Ciclo TCM", "Parada externa PreFeeder"),
    (27, "Ciclo TCM", "Sensores CAN offline (ciclo)"),
    (30, "PreFeeder", "No inicializado"),
    (31, "PreFeeder", "Peer TCP perdido"),
    (32, "Alimentación", "Target/offset inválido"),
    (40, "Lineal", "HOME fallido"),
    (41, "Lineal", "Safe zone fallida"),
    (42, "Lineal", "Movimiento fallido"),
    (43, "Ciclo TCM", "Stop rechazado"),
    (44, "Ciclo TCM", "Reset errores fallido"),
    (45, "PreFeeder", "Disparo TCP fallido"),
    (46, "Alimentación", "Perfil/velocidad inválida (UI)"),
    (47, "Alimentación", "Calibración fallida (UI)"),
    (48, "Alimentación", "Offset fallido (UI)"),
    (49, "Alimentación", "Prueba feed fallida (UI)"),
    (50, "Config", "Depósito no guardado"),
    (51, "Config", "Velocidad no guardada"),
    (52, "Monitoreo", "Cycle time no reseteado"),
    (53, "Modelos", "Operación de modelo fallida"),
    (54, "UI", "Parámetros inválidos"),
    (55, "Ciclo TCM", "Ciclo en conflicto"),
    (56, "PLC", "Seguridad activa"),
    (57, "PreFeeder", "Modo Materialista activo"),
    (58, "Sistema", "Almacenamiento no disponible"),
    (59, "Servo", "ASDA OM no responde"),
    (60, "Alimentación", "Target L + offset negativo"),
    (61, "Alimentación", "Target R + offset negativo"),
    (62, "Alimentación", "Perfil L inválido"),
    (63, "Alimentación", "Perfil R inválido"),
    (64, "Alimentación", "Encoder 6064 L no leído"),
    (65, "Alimentación", "Encoder 6064 R no leído"),
    (66, "Alimentación", "Timeout feed paralelo"),
    (67, "Alimentación", "Reservado (láser timeout)"),
    (68, "Alimentación", "Timeout feed step2"),
    (69, "Alimentación", "Feed incompleto"),
    (70, "Alimentación", "Prefetch/handoff fallido"),
    (71, "Ciclo TCM", "Depósito overview NG"),
    (72, "PreFeeder", "All OK falló en corte"),
    (74, "Alimentación", "Sensor manguera no detectado"),
    (99, "Ciclo TCM", "Falla no clasificada"),
    (100, "Red UI", "Sin conexión con el TCM"),
    (120, "PreFeeder L", "Sin enlace TCP"),
    (121, "PreFeeder L", "Buffer Max (endstop)"),
    (122, "PreFeeder L", "Tensión prolongada"),
    (123, "PreFeeder L", "Cilindro abierto"),
    (124, "PreFeeder L", "Cinta/manguera ausente"),
    (125, "PreFeeder L", "Buffer timeout"),
    (126, "PreFeeder L", "Sin holgura (timeout)"),
    (127, "PreFeeder L", "Parada operador"),
    (130, "PreFeeder R", "Sin enlace TCP"),
    (131, "PreFeeder R", "Buffer Max (endstop)"),
    (132, "PreFeeder R", "Tensión prolongada"),
    (133, "PreFeeder R", "Cilindro abierto"),
    (134, "PreFeeder R", "Cinta/manguera ausente"),
    (135, "PreFeeder R", "Buffer timeout"),
    (136, "PreFeeder R", "Sin holgura (timeout)"),
    (137, "PreFeeder R", "Parada operador"),
]

for code, sys_name, name in UI_CATALOG:
    mod = "TCM / Máquina"
    if sys_name.startswith("PreFeeder"):
        mod = "PreFeeder"
    elif sys_name == "PLC":
        mod = "PLC"
    elif sys_name in ("Lineal", "Servo", "Medición OM", "Alimentación"):
        mod = "Motion / TCM"
    elif sys_name == "Sensores":
        mod = "PLC / Sensores CAN"
    ROWS.append(
        (
            mod,
            "Catálogo UI",
            f"E{code:03d}",
            "",
            "",
            "",
            "",
            "",
            "L" if " L" in sys_name else ("R" if " R" in sys_name else "—"),
            name,
            "",
            "",
            sys_name,
            "Temp/Index.html ERROR_CATALOG",
            "SÍ",
        )
    )

# --- Códigos ciclo firmware (Temp/TCM.ino) ---
CYCLE_FW = [
    ("E000", "OK (no error)", "—"),
    ("E001", "Abort (omitido)", "—"),
    ("E008", "Servo CAN no listo", "E008"),
    ("E010", "Timeout lineal", "E010"),
    ("E011", "Timeout alimentación", "E011"),
    ("E012", "Alimentación fallida", "E012"),
    ("E016", "Sensor manguera alimentación", "E016"),
    ("E020", "Safety stop (genérico)", "E020"),
    ("E021", "Safety pinzas", "E021"),
    ("E022", "Safety sujetador", "E022"),
    ("E023", "Safety cortador", "E023"),
    ("E024", "Safety manguera", "E024"),
    ("E025", "Safety bandeja", "E025"),
    ("E026", "Parada externa PreFeeder", "E026"),
    ("E027", "Sensores CAN offline", "E027"),
    ("E030", "PreFeeder no inicializado", "E030"),
    ("E031", "Peer TCP perdido", "E031"),
    ("E099", "Falla no clasificada", "E099"),
]
for ec, desc, internal in CYCLE_FW:
    if ec == "E000":
        continue  # no es error
    ROWS.append(
        (
            "TCM / Ciclo",
            "Ciclo firmware",
            ec,
            internal,
            "",
            "",
            "",
            "",
            "—",
            desc,
            "setCycleError()",
            "",
            "Ciclo TCM",
            "Temp/TCM.ino cycleErrorCodeLabel",
            "SÍ",
        )
    )

# --- Protocolo Motion ---
MOTION_ERR = [
    (0x0C, 12, "Error (estatus general)", "ErrorState()", "Motion/Motion.ino"),
    (0x11, 17, "Encoder — no detectó cambio en lectura", "Error()", "Motion/Encoder.h"),
    (0x15, 21, "LenghtNG (L) — longitud fuera tolerancia", "LenghtNG_L()", "Motion/FeederCan.h"),
    (0x4B, 75, "LenghtNG (R) — longitud fuera tolerancia", "LenghtNG_R()", "Motion/FeederCan.h"),
]
for hx, dec, name, handler, src in MOTION_ERR:
    ROWS.append(
        (
            "Motion",
            "Protocolo byte TX",
            "",
            "",
            f"0x{hx:02X}",
            dec,
            "",
            "",
            "L" if "(L)" in name else ("R" if "(R)" in name else "—"),
            name,
            handler,
            "",
            "Motion",
            src,
            "SÍ",
        )
    )

# --- Protocolo PLC ---
PLC_ERR = [
    (0x1F, 31, "Cutter Error — sensor cortador", "CutterE()", "PLC/Plc.h"),
    (0x20, 32, "Gripper Error — sensor pinzas", "GripperE()", "PLC/Plc.h"),
    (0x21, 33, "Holder Error — sensor sujetador", "HolderE()", "PLC/Plc.h"),
    (0x22, 34, "Encoder Error — sensor bandeja/encoder", "EncoderE()", "PLC/Plc.h"),
    (0x27, 39, "Error (estatus general PLC)", "ErrorState()", "PLC/Plc.h"),
]
for hx, dec, name, handler, src in PLC_ERR:
    ROWS.append(
        (
            "PLC",
            "Protocolo byte TX",
            "",
            "",
            f"0x{hx:02X}",
            dec,
            "",
            "",
            "—",
            name,
            handler,
            "",
            "PLC",
            src,
            "SÍ",
        )
    )

# --- Protocolo PreFeeder (detalle L/R) ---
# (hex, dec, side, name, handler, wire_l, wire_r)
PF_PROTO = [
    (0x2D, 45, "R", "Buffer Full (R)", "BufferFR()", "", "35"),
    (0x2E, 46, "R", "Buffer Max (R)", "BufferMR()", "", "31"),
    (0x2F, 47, "R", "Tensioner (R)", "TensionerR()", "", "32"),
    (0x30, 48, "R", "Cilindro (R)", "CilindroR()", "", "33"),
    (0x31, 49, "R", "Manguera ausente (R)", "MangueraR()", "", "34"),
    (0x32, 50, "R", "Holgura (R)", "HolguraR()", "", "36"),
    (0x33, 51, "L", "Buffer Full (L)", "BufferFL()", "25", ""),
    (0x34, 52, "L", "Buffer Max (L)", "BufferML()", "21", ""),
    (0x35, 53, "L", "Tensioner (L)", "TensionerL()", "22", ""),
    (0x36, 54, "L", "Cilindro (L)", "CilindroL()", "23", ""),
    (0x37, 55, "L", "Manguera ausente (L)", "MangueraL()", "24", ""),
    (0x38, 56, "L", "Holgura (L)", "HolguraL()", "26", ""),
    (0x3C, 60, "—", "Error (estatus general PreFeeder)", "ErrorState()", "", ""),
]
for hx, dec, side, name, handler, wl, wr in PF_PROTO:
    ROWS.append(
        (
            "PreFeeder",
            "Protocolo byte TX",
            "",
            "",
            f"0x{hx:02X}",
            dec,
            wl,
            wr,
            side,
            name,
            handler,
            "",
            "PreFeeder",
            "PF/PreFeeder_Master/master_tcp.h",
            "SÍ",
        )
    )

# --- PreFeeder lógico PF-001…007 × L/R ---
PF_LOGICAL = [
    ("PF-001", 1, "endstop", "Buffer Max", 3),
    ("PF-002", 2, "tension_timeout", "Tension timeout", 2),
    ("PF-003", 3, "cylinder_open", "Cilindro abierto", 2),
    ("PF-004", 4, "hose_absent", "Manguera ausente", 2),
    ("PF-005", 5, "buffer_timeout", "Buffer sin relleno", 2),
    ("PF-006", 6, "holgura_timeout", "Sin holgura", 2),
    ("PF-007", 7, "operator_stop", "Parada operador", 2),
]
for side, base_wire, base_ui in (("L", 20, 120), ("R", 30, 130)):
    for tag, pid, slug, desc, level in PF_LOGICAL:
        ui = base_ui + pid
        ROWS.append(
            (
                f"PreFeeder {side}",
                "PreFeeder lógico",
                f"E{ui:03d}",
                tag,
                "",
                "",
                str(base_wire + pid) if side == "L" else "",
                str(base_wire + pid) if side == "R" else "",
                side,
                desc,
                slug,
                str(level),
                f"PreFeeder {side}",
                "PF/PF_LR/Status_Mode.h",
                "SÍ",
            )
        )

# --- Máquina / estatus error ---
ROWS.append(
    (
        "TCM / Máquina",
        "Protocolo byte TX",
        "",
        "",
        "0x46",
        70,
        "",
        "",
        "—",
        "Error (estatus general máquina/ciclo)",
        "ErrorState()",
        "",
        "Machine",
        "Doc/protocol_bytes.txt",
        "SÍ",
    )
)

# --- Fases auto con fault PreFeeder ---
for phase in [
    "endstop_fault",
    "tension_fault",
    "cylinder_fault",
    "hose_fault",
    "buffer_fault",
    "holgura_fault",
    "operator_stop",
]:
    ROWS.append(
        (
            "PreFeeder",
            "Fase auto (fault)",
            "",
            "",
            "",
            "",
            "",
            "",
            "L+R",
            f"AutoState: {phase}",
            "pfAutoPhaseName",
            "",
            "PreFeeder",
            "PF/PF_LR/Status_Mode.h PF_AUTO_PHASE_NAMES",
            "SÍ",
        )
    )

# --- Estado máquina PF_MS_ERROR ---
ROWS.append(
    (
        "PreFeeder",
        "Estado operativo",
        "",
        "PF_MS_ERROR",
        "",
        "",
        "",
        "",
        "L+R",
        "Estado operativo: error",
        "pfMachineStateName",
        "",
        "PreFeeder",
        "PF/PF_LR/Status_Mode.h",
        "SÍ",
    )
)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Errores TCM"

    red_fill = PatternFill(start_color="FFCCCB", end_color="FFCCCB", fill_type="solid")
    red_font = Font(color="9C0006", bold=False)
    header_fill = PatternFill(start_color="9C0006", end_color="9C0006", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_data in ROWS:
        ws.append(list(row_data))
        r = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = red_fill
            cell.font = red_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    widths = [14, 16, 10, 12, 10, 8, 8, 8, 6, 42, 22, 6, 18, 28, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Hoja resumen
    summary = wb.create_sheet("Resumen")
    summary.append(["Módulo", "Cantidad"])
    from collections import Counter

    counts = Counter(r[0] for r in ROWS)
    for mod, n in sorted(counts.items()):
        summary.append([mod, n])
    summary.append(["TOTAL", len(ROWS)])
    for row in summary.iter_rows(min_row=1, max_row=summary.max_row):
        for cell in row:
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
            elif cell.column == 1:
                cell.fill = red_fill
                cell.font = red_font

    wb.save(OUT)
    print(f"OK: {OUT} ({len(ROWS)} filas de error)")


if __name__ == "__main__":
    main()
